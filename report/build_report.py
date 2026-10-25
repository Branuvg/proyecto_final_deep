"""Build report/Proyecto_Final.docx from report/sections/*.md using the local template.docx.

Run from the repo root:  python report/build_report.py [--template path/to/template.docx]

Only the text of the cover runs is replaced (formatting is untouched); everything after the cover is
removed and rebuilt with the template's own styles (Heading 1/2, 'Grid Table 1 Light3', Calibri body).
Numbers in section 6 come from outputs/*/results.json, never typed by hand.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

# ---- Cover data (single place to edit) -------------------------------------------------------
TITLE = "Simulador de mundos habitables: predicción de supervivencia con modelos sustitutos y explicabilidad verificable"
COURSE, SECTION, KIND = "Deep Learning", "Sección 10", "Proyecto Final"
MEMBERS = ["David Domínguez, 23712", "Jonathan Díaz, 23837", "Gabriel Bran, 23590"]
REPO = "https://github.com/Branuvg/proyecto_final_deep"
# -----------------------------------------------------------------------------------------------

FONT = "Calibri"


def find_template(arg: str | None) -> Path:
    cands = [Path(arg)] if arg else [ROOT / "template.docx", Path.cwd() / "template.docx"]
    for c in cands:
        if c.exists():
            return c
    sys.exit("ERROR: no se encontró template.docx. Colócalo en la raíz del repositorio "
             f"({ROOT / 'template.docx'}) o indica su ruta con --template. El archivo no se sube al repo (.gitignore).")


def set_run_texts(par, texts: dict[int, str]):
    for i, t in texts.items():
        par.runs[i].text = t


def fill_cover(doc):
    P = doc.paragraphs
    assert P[2].text.startswith("Data") and "30" in P[3].text and "Laboratorio" in P[9].text, "La carátula de la plantilla cambió"
    set_run_texts(P[2], {0: COURSE, 1: ""})
    set_run_texts(P[3], {0: SECTION})
    set_run_texts(P[9], {0: KIND.replace(" ", " "), 1: ""})   # keeps the trailing ':' run
    set_run_texts(P[10], {0: TITLE})
    set_run_texts(P[22], {2: MEMBERS[0] + "  ", 4: MEMBERS[1] + "  ", 6: MEMBERS[2] + " "})
    P[27].runs[-1].text = "Repo: " + REPO


def clear_after_cover(doc):
    body = doc.element.body
    kids = list(body)
    start = next(i for i, k in enumerate(kids) if k.tag == qn("w:p") and "".join(k.itertext()).startswith("1. Introducci"))
    for k in kids[start:]:
        if k.tag != qn("w:sectPr"):
            body.remove(k)


# ---------------------------------------------------------------------------- markdown -> docx
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_inline(par, text: str, size=None, italic=False):
    for tok in INLINE.split(text):
        if not tok:
            continue
        bold = tok.startswith("**") and tok.endswith("**") and len(tok) > 4
        ital = (not bold) and tok.startswith("*") and tok.endswith("*") and len(tok) > 2
        code = tok.startswith("`") and tok.endswith("`")
        r = par.add_run(tok[2:-2] if bold else tok[1:-1] if (ital or code) else tok)
        r.font.name = "Consolas" if code else FONT
        r.bold = bold or None
        r.italic = (ital or italic) or None
        if size:
            r.font.size = Pt(size)


def body_par(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, **kw):
    p = doc.add_paragraph(); p.alignment = align
    add_inline(p, text, **kw)
    return p


def add_table(doc, rows: list[list[str]]):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["Grid Table 1 Light3"]
    pr = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblLook"):
        for e in pr.findall(qn(tag)):
            pr.remove(e)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct"); pr.append(w)
    lk = OxmlElement("w:tblLook")
    for k, v in dict(val="0020", firstRow="1", lastRow="0", firstColumn="0", lastColumn="0", noHBand="0", noVBand="0").items():
        lk.set(qn("w:" + k), v)
    pr.append(lk)
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            p = t.cell(i, j).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(1.8); p.paragraph_format.space_after = Pt(1.8)
            add_inline(p, cell.strip(), size=9 if len(rows[0]) > 4 else None)
    doc.add_paragraph()
    return t


def render(doc, md: str, counters: dict):
    lines = md.splitlines(); i = 0; pending_table_cap = None
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("# "):
            doc.add_paragraph(ln[2:].strip(), style="Heading 1")
        elif ln.startswith("## "):
            doc.add_paragraph(ln[3:].strip(), style="Heading 2")
        elif ln.startswith("[[tabla:"):
            counters["t"] += 1
            cap = ln[len("[[tabla:"):].rstrip("]").strip()
            p = body_par(doc, f"**Tabla {counters['t']}.** {cap}", WD_ALIGN_PARAGRAPH.LEFT, size=10)
            p.paragraph_format.keep_with_next = True
        elif ln.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*)\)", ln)
            path = ROOT / m.group(2)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Inches(6.2 if "cascada" not in path.name else 4.6))
            p.paragraph_format.keep_with_next = True
            counters["f"] += 1
            body_par(doc, f"**Figura {counters['f']}.** {m.group(1)}", WD_ALIGN_PARAGRAPH.CENTER, size=10)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows); continue
        elif ln.startswith("- "):
            p = doc.add_paragraph(style="List Paragraph"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.2)
            add_inline(p, "• " + ln[2:].strip())
        elif ln.startswith("> "):
            body_par(doc, ln[2:].strip(), italic=True)
        elif re.match(r"^\d+\. ", ln):
            body_par(doc, ln)
        elif re.match(r"^\S", ln) and "doi.org" not in ln and len(md) and counters.get("bib"):
            p = body_par(doc, ln, WD_ALIGN_PARAGRAPH.LEFT)       # APA: hanging indent
            p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)
        else:
            p = body_par(doc, ln)
            if counters.get("bib"):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)
        i += 1


# ---------------------------------------------------------------------------- data for section 6
def md_table(rows):
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def build_context():
    import yaml
    from worldsim.data.dataset import load_worlds
    from worldsim.reporting import LABELS, eda, xai_summary
    from worldsim.simulator.factors import FACTORS

    ta = json.load(open(ROOT / "outputs/task_a/results.json"))
    tb = json.load(open(ROOT / "outputs/task_b/results.json"))
    xc = json.load(open(ROOT / "outputs/xai/checks.json"))["checks"]
    ctx = dict(a=ta["results"], n_test=ta["n_test"], b={int(k): v for k, v in tb["results"].items()},
               e=eda(load_worlds(ROOT / "data/worlds.npz")), x=xc)
    xs = xai_summary(str(ROOT / "outputs/xai/attributions.npz"))
    ctx["g"], ctx["act"] = xs["agreement"], xs["actionable_share"]

    rng = yaml.safe_load(open(ROOT / "configs/dataset.yaml", encoding="utf8"))["ranges"]
    fr = [["Factor", "Descripción", "Rango muestreado", "Accionable"]]
    for f in FACTORS:
        lo, hi = rng[f.name]
        fr.append([f"`{f.name}`", f.desc, f"[{lo:g}, {hi:g}]" + (" (log)" if f.log else ""), "Sí" if f.actionable else "No"])
    t_a = [["Modelo", "ROC-AUC", "PR-AUC", "F1", "Brier", "ECE", "MAE vs P verdadera"]]
    for k, v in ta["results"].items():
        t_a.append([LABELS[k], f"{v['roc_auc']:.4f}", f"{v['pr_auc']:.4f}", f"{v['f1']:.4f}", f"{v['brier']:.4f}", f"{v['ece']:.4f}", f"{v['mc_mae']:.4f}"])
    ks = tb["k_percents"]; names = [n for n in LABELS if n in tb["results"][str(ks[0])]]
    t_b = [["k"] + [LABELS[n] for n in names]]
    for k in ks:
        t_b.append([f"{k}%"] + [f"{tb['results'][str(k)][n]['roc_auc']:.4f}" for n in names])
    t_e = [["k", "Planetas test", "Evaluados (vivos)", "Excluidos", "Supervivencia entre vivos"]]
    for k in ks:
        r = tb["results"][str(k)]
        t_e.append([f"{k}%", str(r["n_test"]), str(r["n_test_alive"]), str(r["n_excluded"]), f"{r['pos_rate_alive']:.1%}"])
    imp = xs["importance"]
    t_x = [["Factor"] + list(imp.columns)] + [[f"`{i}`"] + [f"{v:.3f}" for v in row] for i, row in imp.iterrows()]
    ctx["tables"] = dict(factors=md_table(fr), task_a=md_table(t_a), task_b_auc=md_table(t_b), task_b_excl=md_table(t_e), xai_importance=md_table(t_x))
    return ctx


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--template"); ap.add_argument("--out", default=str(ROOT / "report/Proyecto_Final.docx"))
    a = ap.parse_args()
    doc = Document(find_template(a.template))
    fill_cover(doc); clear_after_cover(doc)
    ctx = build_context(); tables = ctx.pop("tables")
    counters = dict(t=0, f=0)
    for f in sorted((ROOT / "report/sections").glob("*.md")):
        md = f.read_text(encoding="utf8")
        md = re.sub(r"<<TABLE (\w+)>>", lambda m: "@@" + m.group(1) + "@@", md)
        md = md.format(**ctx)
        md = re.sub(r"@@(\w+)@@", lambda m: tables[m.group(1)], md)
        counters["bib"] = f.name.startswith("08")
        render(doc, md, counters)
    doc.core_properties.title = TITLE
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(a.out)
    print("OK ->", a.out, f"({counters['f']} figuras, {counters['t']} tablas)")


if __name__ == "__main__":
    main()
